"""Tests de non-régression pour app/extraction.py.

Le premier cas ci-dessous correspond à un bug trouvé chez AI Real-Time
(629a1bb, validation sur ~1500 CV réels) sur le même genre d'extraction
PDF/DOCX que nous portons ici, et observé directement sur le CV utilisé
tout au long de cette fonctionnalité (titre "DÉVELOPPEUR WEB" rendu
"D É V E L O P P E U R W E B" par l'extraction PyMuPDF sur un template au
style décoratif).

Les cas _is_plausible_table/_detect_column_gutter ci-dessous portent deux
bugs réels trouvés par AI Real-Time sur un corpus de validation de 13 715
CV (2026-09-14, fa63df2 et 38d44a8) sur ces mêmes fonctions, portées à
l'identique côté Keoni (point 9).

Le cas test_docx_captures_text_hidden_in_a_textbox porte un bug réel
trouvé par AI Real-Time sur 888 vrais CV DOCX de production (2026-09-18,
50e976a) : le texte placé dans une zone de texte/forme (<w:txbxContent>)
était entièrement invisible à l'extraction -- 165 fichiers (18,6%) avaient
plus de 500 caractères ainsi cachés, certains jusqu'à l'intégralité du CV.
"""

from __future__ import annotations

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.extraction import (
    _collapse_letter_spacing,
    _detect_column_gutter,
    _image_coverage,
    _is_plausible_table,
    clean_text,
    extract_text_from_docx,
    looks_unreadable,
    page_ocr_reason,
)


def test_letter_spaced_heading_is_collapsed():
    assert _collapse_letter_spacing("C O M P É T E N C E S") == "COMPÉTENCES"
    assert _collapse_letter_spacing("D É V E LO P P E U R W E B") == "DÉVELOPPEURWEB"


def test_normal_prose_is_left_alone():
    assert _collapse_letter_spacing("Ceci est une phrase normale") == "Ceci est une phrase normale"
    # Mots d'une lettre légitimes dans une phrase normale (proportion trop faible pour déclencher).
    assert _collapse_letter_spacing("Il a dit à y aller") == "Il a dit à y aller"


def test_short_line_is_not_collapsed():
    # Sous le seuil de 4 tokens : jamais touché, même si tout est en lettres isolées.
    assert _collapse_letter_spacing("R et D") == "R et D"


def test_clean_text_applies_letter_spacing_collapse():
    result = clean_text("D É V E L O P P E U R W E B\nCompétences: Python, Docker")
    assert "DÉVELOPPEURWEB" in result
    assert "Python" in result


def test_whole_page_misdetected_as_a_table_is_rejected():
    """Régression réelle (validation du corpus de 13 715 CV chez AI
    Real-Time, 2026-09-14, CV Rachid Aissaoui) : sur un CV à deux colonnes
    SANS AUCUN cadre de tableau, le détecteur géométrique de PyMuPDF a
    quand même pris la page ENTIÈRE pour un tableau 2 lignes / 2 colonnes
    -- chaque "cellule" étant le texte complet d'une colonne (nom, contact,
    tout un paragraphe...), soit une douzaine de lignes. Le nombre de
    lignes seul ne suffit pas à rejeter ce faux tableau (il en avait bien
    2, comme un petit tableau légitime) -- c'est la taille démesurée de ses
    cellules qui le trahit."""
    rows = [
        [
            "Rachid\nAISSAOUI\nConsultant\nQA\n78280 Guyancourt\n06.24.11.53.98\n"
            "aissaoui.r@gmail.com\nlinkedin.com/in/riln\nPolyvalent et curieux\n"
            "d'adaptation et d'apprentissage\nnouveaux defis\nDIPLOMES / CERTIFICATIONS",
            "EXPERIENCE PROFESSIONNELLE\nCONSULTANT QA\nBig Ben Corp\n"
            "Qualification de sites\nRealisation de tests\nGestion des retours\n"
            "CONSULTANT TEST RECETTE\nGMF Assurances\nProjets Evolutions",
        ],
        ["2013 Licence Assurance\nUniversite Paris II\nISTQB Foundation\nHP Quality Center", None],
    ]
    assert _is_plausible_table(rows) is False


def test_a_real_small_table_is_still_accepted():
    """Garde-fou : un vrai (petit) tableau CV, cellules d'une ou deux
    lignes, ne doit pas être rejeté par le garde-fou de taille."""
    rows = [
        ["Langages", "Python, Java, SAS, SQL, R et Scala"],
        ["Outils de\ndéveloppement", "Toad Siebel SAS6 SAS91 SAS92 SAS94"],
        ["SGBD", "Oracle, PostgreSQL, MySQL, MongoDB"],
    ]
    assert _is_plausible_table(rows) is True


def test_a_single_row_table_is_rejected():
    assert _is_plausible_table([["Langages", "Python, Java"]]) is False


def test_narrow_real_gutter_is_not_rejected_as_pixel_noise():
    """Régression réelle (validation du corpus de 13 715 CV chez AI
    Real-Time, 2026-09-14, CV Aurelien Torres) : le vrai gouffre entre les
    deux colonnes de ce CV ne mesurait que ~20.5pt sur une page de 595pt de
    large (3.4% de la largeur) -- juste sous l'ancien seuil de 4% (23.8pt),
    qui traitait ce gouffre bien réel comme du bruit et retombait sur le
    milieu exact de page (297.5), en plein milieu de la colonne de droite."""
    left = [
        {"bbox": (40.0, 60.0, 113.3, 63.0), "text": "Jean DUPONT"},
        {"bbox": (40.0, 90.0, 103.6, 93.0), "text": "Certifications"},
        {"bbox": (40.0, 120.0, 95.6, 123.0), "text": "Profil perso"},
        {"bbox": (40.0, 150.0, 108.5, 153.0), "text": "Curiosite forte"},
    ]
    right = [
        {"bbox": (134.0, 60.0, 232.4, 63.0), "text": "Chef de Projet MOA"},
        {"bbox": (134.0, 90.0, 472.1, 93.0), "text": "Pilotage de projets SI en environnement multi-pays"},
        {"bbox": (134.0, 120.0, 183.5, 123.0), "text": "Formation"},
        {"bbox": (134.0, 150.0, 258.7, 153.0), "text": "Master informatique 2018"},
    ]
    mid = _detect_column_gutter(left + right, page_width=595.0)
    assert 113.3 < mid < 134.0, (
        f"le vrai gouffre (113.3-134.0) ne doit pas être rejeté au profit "
        f"du milieu de page fixe (297.5), obtenu mid={mid}"
    )


def test_docx_captures_text_hidden_in_a_textbox(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("En-tête visible normalement.")

    # <w:txbxContent> minimal imbriqué dans un run, comme le font les vrais
    # templates de CV pour placer toute une barre latérale/section dans une
    # forme "zone de texte". python-docx n'a pas d'API haut niveau pour les
    # zones de texte, donc ceci est construit directement en XML, de la même
    # façon que python-docx lui-même l'émettrait.
    textbox_xml = (
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml">'
        '<w:pict><v:shape>'
        '<v:textbox><w:txbxContent>'
        '<w:p><w:r><w:t>Compétences: Python, Django, PostgreSQL</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Expérience: 5 ans en développement backend</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox>'
        '</v:shape></w:pict></w:r>'
    )
    p = doc.add_paragraph()
    p._p.append(parse_xml(textbox_xml))
    doc.add_paragraph("Pied de page visible normalement.")

    path = tmp_path / "cv_with_textbox.docx"
    doc.save(path)

    result = extract_text_from_docx(path)

    assert "En-tête visible normalement." in result
    assert "Pied de page visible normalement." in result
    assert "Compétences: Python, Django, PostgreSQL" in result
    assert "Expérience: 5 ans en développement backend" in result


# ── Texte présent mais inexploitable / incomplet : l'OCR doit prendre le relais ──
#
# Constaté en production le 2026-10-12 sur deux CV qui n'avaient AUCUNE
# compétence détectée et remontaient en tête de chaque classement : l'un avait
# une police PDF sans table Unicode (texte décalé), l'autre un corps de CV rendu
# en image avec seulement l'état civil en couche texte (928 caractères). Dans
# les deux cas la page dépassait largement le seuil des « 20 caractères » : l'OCR
# ne se déclenchait pas. Les chaînes ci-dessous ne contiennent aucune donnée
# personnelle.

GARBLED_FONT_TEXT = (
    r"3\WKRQ64/DLQVLTX HQ JHVWLRQGHEDVHVGHGRQQ«HVJ«RJUDSKLTXHVFRPSOH[HVDYHF3RVW*,6RQH[S«ULHQFHPDSHUPLVG "
    r"RSWLPLVHU O H[SORLWDWLRQGHVGRQQ«HVSRXUGHVSURMHWVDQDO\WLTXHVHQJDUDQWLVVDQWODTXDOLW«HWODILDELOLW«GHVLQIRUPDWLRQV"
)

NORMAL_FRENCH_TEXT = (
    "Ingénieur DevOps avec huit ans d'expérience dans la mise en place de pipelines CI/CD et "
    "l'automatisation des déploiements sur des environnements cloud. Responsable de la sécurité des "
    "applications et de la formation des équipes de développement."
)

SKILLS_LIST_TEXT = (
    "Python, SQL, Docker, Kubernetes, Terraform, Ansible, Jenkins, GitLab CI, Prometheus, Grafana, Linux, "
    "Bash, Java, Spring Boot, PostgreSQL, Oracle, MongoDB, Redis, Kafka, RabbitMQ, OpenShift, Azure, AWS, "
    "GCP, SonarQube, Nexus"
)

CIVIL_STATUS_TEXT = (
    "Prénom Nom Né le 19/08/1994 Marocain, Célibataire 28 Rue Jbel Aouam Casablanca Maroc Permis B "
    "Disponible immédiatement Mobilité nationale et internationale Langues Arabe Français Anglais"
)


def test_garbled_font_text_is_detected_as_unreadable():
    assert looks_unreadable(GARBLED_FONT_TEXT) is True


def test_readable_texts_are_not_flagged_as_unreadable():
    assert looks_unreadable(NORMAL_FRENCH_TEXT) is False
    # Une page faite d'une simple liste d'outils n'est pas du bruit.
    assert looks_unreadable(SKILLS_LIST_TEXT) is False
    assert looks_unreadable(CIVIL_STATUS_TEXT) is False


def test_short_text_is_never_called_unreadable():
    assert looks_unreadable("Python SQL") is False
    assert looks_unreadable("") is False


def test_page_ocr_reason_covers_the_three_cases():
    assert page_ocr_reason("", 0.0) == "vide"
    assert page_ocr_reason("quelques mots", 0.0) == "vide"
    assert page_ocr_reason(GARBLED_FONT_TEXT, 0.0) == "illisible"
    # Peu de texte sur une page presque entièrement occupée par une image.
    assert page_ocr_reason(CIVIL_STATUS_TEXT, 0.8) == "incomplète"


def test_page_ocr_reason_leaves_a_normal_text_page_alone():
    assert page_ocr_reason(NORMAL_FRENCH_TEXT, 0.0) is None
    # Une petite photo d'identité sur une page bien remplie ne déclenche rien.
    assert page_ocr_reason(NORMAL_FRENCH_TEXT * 12, 0.1) is None
    # Une grande image mais un texte abondant : le texte suffit.
    assert page_ocr_reason(NORMAL_FRENCH_TEXT * 12, 0.8) is None


def test_image_coverage_measures_how_much_of_the_page_is_an_image():
    import fitz

    text_only = fitz.open()
    page = text_only.new_page(width=600, height=800)
    page.insert_text((50, 100), "Texte seul")
    assert _image_coverage(page) == 0.0

    with_image = fitz.open()
    page = with_image.new_page(width=600, height=800)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100), False)
    pix.set_rect(pix.irect, (200, 200, 200))
    page.insert_image(fitz.Rect(0, 0, 600, 800), pixmap=pix)
    assert _image_coverage(page) >= 0.9


# ── Tesseract limité à un thread (saturation du serveur en OCR parallèle) ─────
#
# Mesuré en production (VPS 4 cœurs) : 8 OCR en parallèle avec les threads par
# défaut de Tesseract (OpenMP) portaient la charge à 32 et dépassaient tous le
# délai de 25 s par page -- les CV illisibles restaient illisibles --, alors que
# le même lot passait en 2,4 s avec OMP_THREAD_LIMIT=1.

def test_single_thread_wrapper_is_used_when_installed(monkeypatch):
    from app import extraction

    monkeypatch.setattr(extraction.shutil, "which", lambda name: "/usr/local/bin/tesseract-1t" if name == "tesseract-1t" else None)
    assert extraction.single_thread_tesseract_cmd() == "/usr/local/bin/tesseract-1t"


def test_no_wrapper_leaves_tesseract_untouched(monkeypatch):
    from app import extraction

    monkeypatch.setattr(extraction.shutil, "which", lambda name: None)
    assert extraction.single_thread_tesseract_cmd() is None


def test_dockerfile_installs_the_wrapper_and_does_not_limit_threads_globally():
    from pathlib import Path

    dockerfile = Path(__file__).resolve().parents[1].joinpath("Dockerfile").read_text(encoding="utf-8")
    assert "tesseract-1t" in dockerfile
    assert "OMP_THREAD_LIMIT=1" in dockerfile
    # jamais en ENV : cela limiterait aussi PyTorch (embeddings, cross-encoder)
    assert not any(
        line.strip().startswith("ENV") and "OMP_THREAD_LIMIT" in line for line in dockerfile.splitlines()
    )
